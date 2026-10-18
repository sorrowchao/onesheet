"""导出 xlsx / csv。

- xlsx：每张 ResultTable 写成一个 sheet，表头加粗、加边框、自动列宽。
- csv：每张表导出一个独立文件（utf-8-sig，Excel 中文不乱码）。
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .models import ResultTable

_INVALID_SHEET_CHARS = set('[]:*?/\\')
_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center")


def _safe_sheet_name(name: str) -> str:
    """清理 Excel 不支持的 sheet 名字符，并截断到 31 字符。"""
    cleaned = "".join("_" if ch in _INVALID_SHEET_CHARS else ch for ch in name)
    return cleaned[:31] or "Sheet"


def _column_widths(rt: ResultTable) -> list[int]:
    widths = []
    for i, col in enumerate(rt.columns):
        max_len = len(str(col))
        for row in rt.rows:
            if i < len(row) and row[i] is not None:
                max_len = max(max_len, len(str(row[i])))
        widths.append(min(max(max_len, 8) + 2, 40))
    return widths


def _write_table(ws, rt: ResultTable) -> None:
    ws.append(rt.columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rt.rows:
        ws.append(list(row))

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.border = _BORDER
            cell.alignment = _CENTER

    for i, w in enumerate(_column_widths(rt), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def export_xlsx(tables: dict[str, ResultTable], path: str) -> None:
    """把多张结果表写进一个 .xlsx（一张表一个 sheet）。"""
    wb = Workbook()
    wb.remove(wb.active)
    if not tables:
        wb.create_sheet("Sheet1")
    for name, rt in tables.items():
        ws = wb.create_sheet(title=_safe_sheet_name(name))
        _write_table(ws, rt)
    wb.save(path)


def export_csv(tables: dict[str, ResultTable], dir_path: str) -> list[str]:
    """每张结果表导出一个 .csv 文件，返回写出的文件路径列表。"""
    out_dir = Path(dir_path)
    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[str] = []
    for name, rt in tables.items():
        fpath = out_dir / (_safe_sheet_name(name) + ".csv")
        with open(fpath, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(rt.columns)
            for row in rt.rows:
                writer.writerow(row)
        written.append(str(fpath))
    return written
