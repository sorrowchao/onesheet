"""writer 单元测试（反向写入回填）。"""
import tempfile
import unittest
from pathlib import Path

import xlwt
from openpyxl import Workbook

from app.core.models import (FieldMapping, OutputTable, ResultTable, SheetConfig)
from app.core.reader import read_excel
from app.core.writer import (_norm_key, preview_write_back,
                             read_summary_xlsx, write_back)


def make_output_table():
    return OutputTable(
        name="汇总",
        sheet=SheetConfig(sheet_keyword="工资表"),
        fields=[
            FieldMapping(target="姓名", source="kv", key="姓名", value_row=1, value_col=2),
            FieldMapping(target="性别", source="kv", key="性别", value_row=2, value_col=2),
            FieldMapping(target="年龄", source="kv", key="年龄", value_row=3, value_col=2),
        ],
    )


def make_form_xlsx(path: str, name, gender, age):
    wb = Workbook()
    ws = wb.active
    ws.title = "工资表"
    ws.cell(row=1, column=1).value = "姓名"
    ws.cell(row=1, column=2).value = name
    ws.cell(row=2, column=1).value = "性别"
    ws.cell(row=2, column=2).value = gender
    ws.cell(row=3, column=1).value = "年龄"
    ws.cell(row=3, column=2).value = age
    wb.save(path)
    wb.close()


def make_form_xls(path: str, name, gender, age):
    wb = xlwt.Workbook()
    ws = wb.add_sheet("工资表")
    ws.write(0, 0, "姓名")
    ws.write(0, 1, name)
    ws.write(1, 0, "性别")
    ws.write(1, 1, gender)
    ws.write(2, 0, "年龄")
    ws.write(2, 1, age)
    wb.save(path)


def cell_value(path: str, row: int, col: int):
    """读回写出的文件第 row 行第 col 列（0-based）。"""
    sheet = read_excel(path)[0]
    return sheet.grid[row][col]


class TestWriteBackXlsx(unittest.TestCase):
    def test_round_trip(self):
        with tempfile.TemporaryDirectory() as tmp:
            p1 = str(Path(tmp) / "张三.xlsx")
            p2 = str(Path(tmp) / "李四.xlsx")
            make_form_xlsx(p1, "张三", "男", 34)
            make_form_xlsx(p2, "李四", "女", 25)

            summary = ResultTable(
                name="汇总",
                columns=["姓名", "性别", "年龄"],
                rows=[["张三", "女", 35], ["李四", "男", 26]],
            )
            out_dir = str(Path(tmp) / "out")
            report = write_back([p1, p2], make_output_table(), summary, [0], out_dir)

            self.assertEqual(report["unmatched"], [])
            self.assertEqual(report["errors"], [])
            self.assertEqual(report["unused_rows"], [])
            self.assertEqual(len(report["written"]), 2)

            o1 = str(Path(out_dir) / "张三.xlsx")
            o2 = str(Path(out_dir) / "李四.xlsx")
            self.assertEqual(cell_value(o1, 1, 1), "女")
            self.assertEqual(cell_value(o1, 2, 1), 35)
            self.assertEqual(cell_value(o2, 1, 1), "男")
            self.assertEqual(cell_value(o2, 2, 1), 26)

    def test_unmatched_and_unused(self):
        with tempfile.TemporaryDirectory() as tmp:
            p1 = str(Path(tmp) / "张三.xlsx")
            make_form_xlsx(p1, "张三", "男", 34)

            # 汇总里只有「王五」，与文件「张三」不匹配
            summary = ResultTable(
                name="汇总", columns=["姓名", "性别", "年龄"],
                rows=[["王五", "男", 30]],
            )
            report = write_back([p1], make_output_table(), summary, [0], str(Path(tmp) / "out"))

            self.assertEqual(report["written"], [])
            self.assertEqual(report["unmatched"], [p1])
            self.assertEqual(report["errors"], [])
            self.assertEqual(report["unused_rows"], [1])

    def test_none_skipped(self):
        with tempfile.TemporaryDirectory() as tmp:
            p1 = str(Path(tmp) / "张三.xlsx")
            make_form_xlsx(p1, "张三", "男", 34)

            # 年龄为 None：不写回，原值保留；性别正常写回
            summary = ResultTable(
                name="汇总", columns=["姓名", "性别", "年龄"],
                rows=[["张三", "女", None]],
            )
            out_dir = str(Path(tmp) / "out")
            report = write_back([p1], make_output_table(), summary, [0], out_dir)

            self.assertEqual(report["errors"], [])
            o1 = str(Path(out_dir) / "张三.xlsx")
            self.assertEqual(cell_value(o1, 1, 1), "女")
            self.assertEqual(cell_value(o1, 2, 1), 34)  # None 未覆盖原值

    def test_composite_key_disambiguates(self):
        with tempfile.TemporaryDirectory() as tmp:
            p1 = str(Path(tmp) / "张三A.xlsx")
            p2 = str(Path(tmp) / "张三B.xlsx")
            make_form_xlsx(p1, "张三", "男", 34)
            make_form_xlsx(p2, "张三", "女", 25)

            # 两个同名「张三」，仅靠姓名会重复匹配；用 [姓名, 性别] 复合区分
            summary = ResultTable(
                name="汇总", columns=["姓名", "性别", "年龄"],
                rows=[["张三", "女", 35], ["张三", "男", 40]],
            )
            out_dir = str(Path(tmp) / "out")
            report = write_back([p1, p2], make_output_table(), summary, [0, 1], out_dir)

            self.assertEqual(report["unmatched"], [])
            self.assertEqual(report["errors"], [])
            self.assertEqual(len(report["written"]), 2)
            o1 = str(Path(out_dir) / "张三A.xlsx")
            o2 = str(Path(out_dir) / "张三B.xlsx")
            self.assertEqual(cell_value(o1, 2, 1), 40)  # 张三A(男) → 年龄 40
            self.assertEqual(cell_value(o2, 2, 1), 35)  # 张三B(女) → 年龄 35

    def test_written_cell_centered_other_format_preserved(self):
        import openpyxl
        from openpyxl.styles import Alignment, Font

        with tempfile.TemporaryDirectory() as tmp:
            p = str(Path(tmp) / "张三.xlsx")
            wb = Workbook()
            ws = wb.active
            ws.title = "工资表"
            ws.cell(row=1, column=1).value = "姓名"
            ws.cell(row=1, column=2).value = "张三"
            ws.cell(row=2, column=1).value = "性别"
            ws.cell(row=2, column=2).value = "男"
            ws.cell(row=1, column=2).font = Font(bold=True)
            ws.cell(row=1, column=2).alignment = Alignment(horizontal="left")
            ws.cell(row=9, column=9).value = "不该动"
            ws.cell(row=9, column=9).font = Font(color="FFCC0000", size=20)
            wb.save(p)
            wb.close()

            summary = ResultTable(
                name="汇总", columns=["姓名", "性别"], rows=[["张三", "女"]])
            out_dir = str(Path(tmp) / "out")
            report = write_back([p], make_output_table(), summary, [0], out_dir)
            self.assertEqual(report["errors"], [])

            o = openpyxl.load_workbook(str(Path(out_dir) / "张三.xlsx"))["工资表"]
            # 回填单元格：默认水平居中，原有加粗保留
            self.assertEqual(o.cell(row=1, column=2).alignment.horizontal, "center")
            self.assertTrue(o.cell(row=1, column=2).font.bold)
            self.assertEqual(o.cell(row=1, column=2).value, "张三")
            # 未回填的单元格原格式不变
            self.assertEqual(o.cell(row=9, column=9).value, "不该动")
            self.assertEqual(o.cell(row=9, column=9).font.color.rgb, "FFCC0000")
            o.parent.close()

    def test_field_order_does_not_affect_write_back(self):
        """字段顺序打乱后，回填仍按 target/取值位置匹配，不受列表顺序影响。"""
        with tempfile.TemporaryDirectory() as tmp:
            p1 = str(Path(tmp) / "张三.xlsx")
            make_form_xlsx(p1, "张三", "男", 34)

            ot = make_output_table()
            ot.fields.reverse()  # 打乱字段顺序

            summary = ResultTable(
                name="汇总", columns=["姓名", "性别", "年龄"],
                rows=[["张三", "女", 35]],
            )
            out_dir = str(Path(tmp) / "out")
            report = write_back([p1], ot, summary, [0], out_dir)

            self.assertEqual(report["errors"], [])
            self.assertEqual(report["unmatched"], [])
            o1 = str(Path(out_dir) / "张三.xlsx")
            self.assertEqual(cell_value(o1, 1, 1), "女")
            self.assertEqual(cell_value(o1, 2, 1), 35)

    def test_missing_key_field(self):
        with tempfile.TemporaryDirectory() as tmp:
            p1 = str(Path(tmp) / "张三.xlsx")
            make_form_xlsx(p1, "张三", "男", 34)

            # 匹配列「工号」在模板字段里没有定位位置
            summary = ResultTable(
                name="汇总", columns=["工号", "姓名"],
                rows=[["001", "张三"]],
            )
            report = write_back([p1], make_output_table(), summary, [0], str(Path(tmp) / "out"))

            self.assertEqual(report["written"], [])
            self.assertEqual(report["unmatched"], [])
            self.assertEqual(len(report["errors"]), 1)
            self.assertIn("工号", report["errors"][0][1])


class TestWriteBackXls(unittest.TestCase):
    def test_xls_round_trip(self):
        with tempfile.TemporaryDirectory() as tmp:
            p1 = str(Path(tmp) / "张三.xls")
            make_form_xls(p1, "张三", "男", 34)

            summary = ResultTable(
                name="汇总", columns=["姓名", "性别", "年龄"],
                rows=[["张三", "女", 35]],
            )
            out_dir = str(Path(tmp) / "out")
            report = write_back([p1], make_output_table(), summary, [0], out_dir)

            self.assertEqual(report["errors"], [])
            self.assertEqual(report["unmatched"], [])
            self.assertEqual(len(report["written"]), 1)
            o1 = str(Path(out_dir) / "张三.xls")
            self.assertEqual(cell_value(o1, 1, 1), "女")
            self.assertEqual(cell_value(o1, 2, 1), 35)

    def test_xls_written_cell_keeps_format_and_center(self):
        import xlrd

        with tempfile.TemporaryDirectory() as tmp:
            p = str(Path(tmp) / "张三.xls")
            wb = xlwt.Workbook()
            ws = wb.add_sheet("工资表")
            styled = xlwt.easyxf("font: bold on; align: horiz left; borders: top thin")
            ws.write(0, 0, "姓名")
            ws.write(0, 1, "张三", styled)
            ws.write(1, 0, "性别")
            ws.write(1, 1, "男", styled)
            ws.write(2, 0, "年龄")
            ws.write(2, 1, 34)
            wb.save(p)

            summary = ResultTable(
                name="汇总", columns=["姓名", "性别"], rows=[["张三", "女"]])
            out_dir = str(Path(tmp) / "out")
            report = write_back([p], make_output_table(), summary, [0], out_dir)
            self.assertEqual(report["errors"], [])

            rb = xlrd.open_workbook(str(Path(out_dir) / "张三.xls"), formatting_info=True)
            s = rb.sheet_by_index(0)
            self.assertEqual(s.cell_value(1, 1), "女")            # 性别回填
            xf = rb.xf_list[s.cell_xf_index(1, 1)]
            self.assertTrue(rb.font_list[xf.font_index].bold)     # 原加粗保留
            self.assertEqual(xf.alignment.hor_align, 2)          # 居中
            self.assertEqual(xf.border.top_line_style, 1)        # 原边框保留


class TestPreviewWriteBack(unittest.TestCase):
    def test_preview_counts_without_writing(self):
        with tempfile.TemporaryDirectory() as tmp:
            p1 = str(Path(tmp) / "张三.xlsx")
            p2 = str(Path(tmp) / "李四.xlsx")
            make_form_xlsx(p1, "张三", "男", 34)
            make_form_xlsx(p2, "李四", "女", 25)
            summary = ResultTable(
                name="汇总", columns=["姓名", "性别", "年龄"],
                rows=[["张三", "女", 35], ["李四", "男", 26]],
            )
            out_dir = Path(tmp) / "out"
            plan = preview_write_back([p1, p2], make_output_table(), summary, [0])

            self.assertEqual(plan["file_count"], 2)
            self.assertEqual(plan["cell_count"], 6)   # 2 文件 × 3 字段
            self.assertEqual(plan["unmatched"], [])
            self.assertEqual(plan["errors"], [])
            self.assertEqual(plan["unused_rows"], [])
            self.assertFalse(plan["cancelled"])
            self.assertFalse(out_dir.exists())        # 预演不写任何东西

    def test_preview_unmatched(self):
        with tempfile.TemporaryDirectory() as tmp:
            p1 = str(Path(tmp) / "张三.xlsx")
            make_form_xlsx(p1, "张三", "男", 34)
            summary = ResultTable(
                name="汇总", columns=["姓名", "性别", "年龄"],
                rows=[["王五", "男", 30]],
            )
            plan = preview_write_back([p1], make_output_table(), summary, [0])

            self.assertEqual(plan["file_count"], 0)
            self.assertEqual(plan["cell_count"], 0)
            self.assertEqual(plan["unmatched"], [p1])
            self.assertEqual(plan["unused_rows"], [1])


class TestCancel(unittest.TestCase):
    def test_cancel_before_start_writes_nothing(self):
        with tempfile.TemporaryDirectory() as tmp:
            p1 = str(Path(tmp) / "张三.xlsx")
            make_form_xlsx(p1, "张三", "男", 34)
            summary = ResultTable(
                name="汇总", columns=["姓名", "性别", "年龄"],
                rows=[["张三", "女", 35]],
            )
            out_dir = str(Path(tmp) / "out")
            report = write_back([p1], make_output_table(), summary, [0], out_dir,
                                cancel_cb=lambda: True)

            self.assertTrue(report["cancelled"])
            self.assertEqual(report["written"], [])
            self.assertEqual(report["unused_rows"], [1])
            self.assertFalse((Path(out_dir) / "张三.xlsx").exists())

    def test_cancel_keeps_files_already_written(self):
        with tempfile.TemporaryDirectory() as tmp:
            p1 = str(Path(tmp) / "张三.xlsx")
            p2 = str(Path(tmp) / "李四.xlsx")
            make_form_xlsx(p1, "张三", "男", 34)
            make_form_xlsx(p2, "李四", "女", 25)
            summary = ResultTable(
                name="汇总", columns=["姓名", "性别", "年龄"],
                rows=[["张三", "女", 35], ["李四", "男", 26]],
            )
            out_dir = str(Path(tmp) / "out")
            checks = {"n": 0}

            def cancel():
                checks["n"] += 1
                return checks["n"] > 1   # 处理完第一个文件后取消

            report = write_back([p1, p2], make_output_table(), summary, [0],
                                out_dir, cancel_cb=cancel)

            self.assertTrue(report["cancelled"])
            self.assertEqual(len(report["written"]), 1)
            self.assertEqual(report["unused_rows"], [2])
            self.assertTrue((Path(out_dir) / "张三.xlsx").exists())
            self.assertFalse((Path(out_dir) / "李四.xlsx").exists())


class TestNormKey(unittest.TestCase):
    """键值归一化：跨类型/格式比对应视为同一值。"""

    def test_numeric_and_formatted_strings_match(self):
        self.assertEqual(_norm_key("1,234"), _norm_key(1234))
        self.assertEqual(_norm_key("1,234"), _norm_key(1234.0))
        self.assertEqual(_norm_key("34.0"), _norm_key(34))
        self.assertEqual(_norm_key("34"), _norm_key(34.0))
        self.assertEqual(_norm_key("2025"), _norm_key(2025))
        self.assertEqual(_norm_key("3.5"), _norm_key(3.5))

    def test_leading_zero_text_preserved(self):
        self.assertEqual(_norm_key("007"), "007")
        self.assertNotEqual(_norm_key("007"), _norm_key(7))

    def test_whitespace_and_empty(self):
        self.assertEqual(_norm_key("  张三  "), _norm_key("张三"))
        self.assertEqual(_norm_key(None), "")
        self.assertEqual(_norm_key(""), "")
        self.assertEqual(_norm_key("   "), "")

    def test_write_back_matches_formatted_number_key(self):
        """分表键为文本 "1,234"，汇总键为数值 1234 时也能匹配回填。"""
        with tempfile.TemporaryDirectory() as tmp:
            p1 = str(Path(tmp) / "张三.xlsx")
            wb = Workbook()
            ws = wb.active
            ws.title = "工资表"
            ws.cell(row=1, column=1).value = "工号"
            ws.cell(row=1, column=2).value = "1,234"
            ws.cell(row=2, column=1).value = "姓名"
            ws.cell(row=2, column=2).value = "张三"
            ws.cell(row=3, column=1).value = "性别"
            ws.cell(row=3, column=2).value = "男"
            wb.save(p1)
            wb.close()

            ot = OutputTable(
                name="汇总",
                sheet=SheetConfig(sheet_keyword="工资表"),
                fields=[
                    FieldMapping(target="工号", source="kv", key="工号", value_row=1, value_col=2),
                    FieldMapping(target="姓名", source="kv", key="姓名", value_row=2, value_col=2),
                    FieldMapping(target="性别", source="kv", key="性别", value_row=3, value_col=2),
                ],
            )
            summary = ResultTable(
                name="汇总", columns=["工号", "姓名", "性别"],
                rows=[[1234, "张三", "女"]],
            )
            out_dir = str(Path(tmp) / "out")
            report = write_back([p1], ot, summary, [0], out_dir)

            self.assertEqual(report["unmatched"], [])
            self.assertEqual(report["errors"], [])
            o1 = str(Path(out_dir) / "张三.xlsx")
            self.assertEqual(cell_value(o1, 2, 1), "女")


class TestReadSummary(unittest.TestCase):
    def test_read_summary_xlsx(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "汇总.xlsx")
            wb = Workbook()
            ws = wb.active
            ws.title = "汇总"
            ws.append(["姓名", "性别", "年龄"])
            ws.append(["张三", "男", 30])
            ws.append(["李四", "女", 25])
            wb.save(path)
            wb.close()

            rt = read_summary_xlsx(path)
            self.assertEqual(rt.name, "汇总")
            self.assertEqual(rt.columns, ["姓名", "性别", "年龄"])
            self.assertEqual(rt.rows, [["张三", "男", 30], ["李四", "女", 25]])


if __name__ == "__main__":
    unittest.main()
