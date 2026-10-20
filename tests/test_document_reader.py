"""document_reader / csv_reader / docx_reader / 格式识别单元测试（纯逻辑）。"""
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from app.core.csv_reader import CsvReadError, read_csv
from app.core.document_reader import (is_excel_file, partition_excel_files,
                                      read_document)
from app.core.docx_reader import DocxReadError, read_docx
from app.core.models import Region
from app.core.reader import detect_format, read_excel


def _write(path: Path, data: bytes) -> str:
    path.write_bytes(data)
    return str(path)


def _make_xlsx(path: Path) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "姓名"
    ws["B1"] = "张三"
    wb.save(path)
    wb.close()
    return str(path)


def _make_docx(path: Path) -> str:
    doc = Document()
    doc.add_paragraph("姓名：张三")
    doc.add_paragraph("性别：男")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "单位名称"
    table.cell(0, 1).text = "X企业"
    table.cell(1, 0).text = "项目名称"
    table.cell(1, 1).text = "测试项目"
    doc.save(path)
    return str(path)


class TempDirTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()


class DetectFormatTest(TempDirTest):
    def test_docx_and_xlsx_are_distinguished_inside_zip(self):
        self.assertEqual(detect_format(_make_docx(self.root / "a.docx")), "docx")
        self.assertEqual(detect_format(_make_xlsx(self.root / "a.xlsx")), "xlsx")

    def test_pdf_magic(self):
        path = _write(self.root / "a.pdf", b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
        self.assertEqual(detect_format(path), "pdf")

    def test_csv_by_extension(self):
        path = _write(self.root / "a.csv", "姓名,值\n张三,1\n".encode("utf-8"))
        self.assertEqual(detect_format(path), "csv")

    def test_xls_by_magic(self):
        import xlwt
        wb = xlwt.Workbook()
        wb.add_sheet("s").write(0, 0, "x")
        path = str(self.root / "a.xls")
        wb.save(path)
        self.assertEqual(detect_format(path), "xls")

    def test_old_doc_rejected(self):
        path = _write(self.root / "old.doc",
                      b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32)
        with self.assertRaises(ValueError) as ctx:
            detect_format(path)
        self.assertIn("不支持旧版", str(ctx.exception))

    def test_unknown_extension_rejected(self):
        path = _write(self.root / "a.bin", b"hello")
        with self.assertRaises(ValueError):
            detect_format(path)

    def test_read_excel_rejects_non_excel(self):
        path = _write(self.root / "a.csv", b"a,b\n1,2\n")
        with self.assertRaises(ValueError) as ctx:
            read_excel(path)
        self.assertIn("不是 Excel", str(ctx.exception))


class CsvReaderTest(TempDirTest):
    def test_utf8_bom_chinese(self):
        path = _write(self.root / "a.csv",
                      "姓名,单位\n张三,X企业\n".encode("utf-8-sig"))
        sheets = read_csv(path)
        self.assertEqual(len(sheets), 1)
        self.assertEqual(sheets[0].name, "CSV")
        self.assertEqual(sheets[0].grid, [["姓名", "单位"], ["张三", "X企业"]])

    def test_gb18030(self):
        path = _write(self.root / "a.csv", "姓名,值\n张三,1\n".encode("gb18030"))
        self.assertEqual(read_csv(path)[0].grid[1], ["张三", "1"])

    def test_utf16_with_bom(self):
        path = _write(self.root / "a.csv", "姓名,值\n张三,1\n".encode("utf-16"))
        self.assertEqual(read_csv(path)[0].grid[1], ["张三", "1"])

    def test_tab_and_semicolon_delimiters(self):
        tab = _write(self.root / "t.csv", b"a\tb\tc\n1\t2\t3\n")
        semi = _write(self.root / "s.csv", b"a;b;c\n1;2;3\n")
        self.assertEqual(read_csv(tab)[0].grid[1], ["1", "2", "3"])
        self.assertEqual(read_csv(semi)[0].grid[1], ["1", "2", "3"])

    def test_leading_zeros_preserved(self):
        path = _write(self.root / "z.csv", b"code\n007\n")
        self.assertEqual(read_csv(path)[0].grid[1][0], "007")

    def test_trailing_empty_rows_and_columns_trimmed(self):
        path = _write(self.root / "e.csv", b"a,b,,\n1,2,,\n,,,\n")
        sheet = read_csv(path)[0]
        self.assertEqual(sheet.grid, [["a", "b"], ["1", "2"]])
        self.assertEqual((sheet.max_row, sheet.max_col), (2, 2))

    def test_bad_encoding_raises(self):
        # UTF-16 BOM + 奇数个字节 → 无法解码
        path = _write(self.root / "bad.csv", b"\xff\xfe\x41\x00\x42")
        with self.assertRaises(CsvReadError):
            read_csv(path)


class DocxReaderTest(TempDirTest):
    def test_tables_and_paragraphs(self):
        sheets = read_docx(_make_docx(self.root / "a.docx"))
        self.assertEqual([s.name for s in sheets], ["表1", "正文"])
        self.assertEqual(sheets[0].grid, [["单位名称", "X企业"],
                                          ["项目名称", "测试项目"]])
        self.assertEqual(sheets[1].grid, [["姓名", "张三"], ["性别", "男"]])

    def test_horizontal_merge(self):
        path = self.root / "m.docx"
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "合并"
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(1, 0).text = "A"
        table.cell(1, 1).text = "B"
        doc.save(path)
        sheet = read_docx(str(path))[0]
        self.assertEqual(sheet.grid[0], ["合并", None])
        self.assertEqual(sheet.merged, [Region(1, 1, 1, 2)])

    def test_vertical_merge(self):
        path = self.root / "v.docx"
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "纵向"
        table.cell(0, 0).merge(table.cell(1, 0))
        table.cell(0, 1).text = "X"
        table.cell(1, 1).text = "Y"
        doc.save(path)
        sheet = read_docx(str(path))[0]
        self.assertEqual(sheet.grid[0][0], "纵向")
        self.assertIsNone(sheet.grid[1][0])
        self.assertIn((1, 1, 2, 1), [(r.top, r.left, r.bottom, r.right)
                                     for r in sheet.merged])

    def test_paragraph_without_separator_is_single_column(self):
        path = self.root / "p.docx"
        doc = Document()
        doc.add_paragraph("只是一段普通文字")
        doc.save(path)
        sheets = read_docx(str(path))
        self.assertEqual(sheets[0].name, "正文")
        self.assertEqual(sheets[0].grid, [["只是一段普通文字"]])

    def test_corrupt_docx_raises_friendly_error(self):
        path = _write(self.root / "bad.docx", b"not a docx")
        with self.assertRaises(DocxReadError) as ctx:
            read_docx(path)
        self.assertIn("Word", str(ctx.exception))


class DocumentDispatcherTest(TempDirTest):
    def test_read_document_dispatch(self):
        self.assertEqual(read_document(_make_xlsx(self.root / "a.xlsx"))[0].name,
                         "Sheet1")
        self.assertEqual(read_document(_make_docx(self.root / "a.docx"))[0].name,
                         "表1")
        csv_path = _write(self.root / "a.csv", b"a,b\n1,2\n")
        self.assertEqual(read_document(csv_path)[0].name, "CSV")

    def test_partition_excel_files(self):
        xlsx = _make_xlsx(self.root / "a.xlsx")
        docx = _make_docx(self.root / "a.docx")
        csv = _write(self.root / "a.csv", b"a,b\n")
        pdf = _write(self.root / "a.pdf", b"%PDF-1.4\n")
        excel, skipped = partition_excel_files([docx, xlsx, csv, pdf])
        self.assertEqual(excel, [xlsx])
        self.assertEqual(skipped, [docx, csv, pdf])

    def test_is_excel_file(self):
        self.assertTrue(is_excel_file(_make_xlsx(self.root / "a.xlsx")))
        self.assertFalse(is_excel_file(_write(self.root / "a.csv", b"a,b\n")))
        self.assertFalse(is_excel_file(str(self.root / "missing.xlsx")))


if __name__ == "__main__":
    unittest.main()
